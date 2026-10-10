"""docx_helpers.py - shared Word formatting for the V2 report and explainer:
Times New Roman 11 body, navy headings, indented bullets, plain hyphens and
straight quotes only (via textclean.clean), static contents page with page
numbers measured by src/make_report.py."""
import json
import os
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(__file__))
from textclean import clean

# --------------------------------------------------------------------------- #
# document scaffolding
# --------------------------------------------------------------------------- #
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
NAVYC = RGBColor(0x1F, 0x38, 0x64)
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.08
for h, sz in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
    hs = doc.styles[h]
    hs.font.name = "Times New Roman"; hs.font.size = Pt(sz); hs.font.bold = True
    hs.font.color.rgb = NAVYC
    rpr = hs.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    hs.paragraph_format.space_before = Pt(12 if h == "Heading 1" else 8)
    hs.paragraph_format.space_after = Pt(4)
lb = doc.styles["List Bullet"]
lb.font.name = "Times New Roman"; lb.font.size = Pt(11)
lb.paragraph_format.left_indent = Inches(0.45)
lb.paragraph_format.first_line_indent = Inches(-0.2)
lb.paragraph_format.space_after = Pt(3)


def para(text, italic=False, bold=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(clean(text)); r.italic = italic; r.bold = bold
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(parts, space_after=6):
    """parts: list of (text, bold) tuples in one paragraph."""
    p = doc.add_paragraph()
    for t, b in parts:
        r = p.add_run(clean(t)); r.bold = b
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        r = p.add_run(clean(lead)); r.bold = True
    p.add_run(clean(text))
    return p


def heading(text, level):
    if text != "Contents" and level <= 2:
        HEADS.append((level, clean(text)))
    return doc.add_heading(clean(text), level=level)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)


def table(headers, rows, widths=None, bold_rows=(), caption=None, note=None, size=9):
    if caption:
        c = doc.add_paragraph(); r = c.add_run(clean(caption))
        r.bold = True; r.font.size = Pt(10)
        c.paragraph_format.space_after = Pt(3); c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]; cell.text = ""
        run = cell.paragraphs[0].add_run(clean(h)); run.bold = True; run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        shade(cell, "1F3864")
    for i, rw in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(rw):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(clean(str(val))); run.font.size = Pt(size)
            if i in bold_rows:
                run.bold = True; shade(cells[j], "DCE6F1")
            if j > 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[j] * 1440)))
        for rw in t.rows:
            for j, w in enumerate(widths):
                rw.cells[j].width = Inches(w)
    if note:
        n = doc.add_paragraph(); r = n.add_run(clean(note)); r.italic = True; r.font.size = Pt(9)
        n.paragraph_format.space_after = Pt(8)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


fig_no = [0]


def figure(path, caption, width=6.3):
    fig_no[0] += 1
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    c = doc.add_paragraph(); r = c.add_run(clean(f"Figure {fig_no[0]}. {caption}"))
    r.italic = True; r.font.size = Pt(9.5)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)


tab_no = [0]


def tcap(text):
    tab_no[0] += 1
    return f"Table {tab_no[0]}. {text}"


TOC_FILE = os.environ.get("DOC_TOC", "data/research/_toc.json")
HEADS = []          # (level, text) collected while building, for the next pass


def add_toc():
    """Static contents list with page numbers measured from the previous
    rendering pass (see src/make_report.py)."""
    if not os.path.exists(TOC_FILE):
        return
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    for lvl, text, page in json.load(open(TOC_FILE)):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(1 if lvl == 2 else 3); pf.space_before = Pt(0 if lvl == 2 else 3)
        pf.left_indent = Inches(0.3 if lvl == 2 else 0)
        pf.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(f"{text}\t{page}"); r.font.size = Pt(10 if lvl == 2 else 10.5); r.bold = (lvl == 1)


def add_page_numbers():
    for sec in doc.sections:
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        for tag, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
            if tag:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
            r._r.append(el)
        r.font.size = Pt(9)




def finish(path, heads_file):
    add_page_numbers()
    json.dump(HEADS, open(heads_file, "w"))
    doc.save(path)
