"""
build_excel.py - package every dataset and result used by the project into a
single formatted workbook: crypto_stat_arb_data.xlsx

All figures are computed in Python (src/research.py, src/statarb.py) and
written as values, so the workbook always matches the report and the deck.
The hourly price file (~670k rows) is kept as data/crypto_hourly.csv rather
than embedded.
"""
import json
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from textclean import clean

FONT = "Arial"
NAVY, LIGHT, WHITE, GREYTXT = "1F3864", "DCE6F1", "FFFFFF", "595959"
hdr_font = Font(name=FONT, bold=True, color=WHITE, size=10)
hdr_fill = PatternFill("solid", fgColor=NAVY)
hl_fill = PatternFill("solid", fgColor=LIGHT)
title_font = Font(name=FONT, bold=True, size=15, color=NAVY)
sub_font = Font(name=FONT, italic=True, size=10, color=GREYTXT)
base_font = Font(name=FONT, size=10)
bold_font = Font(name=FONT, size=10, bold=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

R = "data/research"
S = json.load(open(f"{R}/summary.json"))
wb = Workbook()


def fmt_for(col):
    c = col.lower()
    if "date" in c:
        return "yyyy-mm-dd"
    if any(k in c for k in ["sharpe", "sortino", "calmar", "t-stat", "alpha t", "beta", "corr",
                            "net20", "net7", "gross is", "gross oos", "gross full", "selection", "net sharpe"]):
        return "0.00"
    if any(k in c for k in ["return", "vol", "maxdd", "max dd", "cagr", "alpha", "hitrate", "weight",
                            "drag", "pct_of", "drawdown", "ann.ret"]):
        return "0.0%"
    if "bps" in c:
        return "0.0"
    if "turnover" in c:
        return "0.000"
    return None


def sheet(name, df, title, subtitle, index=False, widths=None, highlight=None, first_w=34, fmt_override=None):
    ws = wb.create_sheet(clean(name)[:31])
    ws["A1"] = clean(title); ws["A1"].font = title_font
    ws["A2"] = clean(subtitle); ws["A2"].font = sub_font
    if index:
        df = df.reset_index()
    df = df.copy()
    df.columns = [clean(str(c)) for c in df.columns]
    r0 = 4
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=r0, column=j, value=col)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
    fmts = {col: (fmt_override or {}).get(col, fmt_for(col)) for col in df.columns}
    vals = df.values
    for i in range(len(df)):
        hl = highlight is not None and highlight(df.iloc[i])
        for j, col in enumerate(df.columns, 1):
            v = vals[i, j - 1]
            if isinstance(v, (float, np.floating)) and (np.isnan(v) or np.isinf(v)):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            elif isinstance(v, str):
                v = clean(v)
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            c = ws.cell(row=r0 + 1 + i, column=j, value=v)
            c.font = bold_font if hl else base_font
            c.border = border
            if hl:
                c.fill = hl_fill
            if fmts[col] and not isinstance(v, str):
                c.number_format = fmts[col]
            c.alignment = left if j == 1 else Alignment(horizontal="center")
    ws.column_dimensions["A"].width = first_w
    for j in range(2, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 14)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    return ws


def big_sheet(name, df, title, subtitle, fmt="0.0000%"):
    """Fast writer for long daily tables (no per-cell borders)."""
    ws = wb.create_sheet(name[:31])
    ws["A1"] = clean(title); ws["A1"].font = title_font
    ws["A2"] = clean(subtitle); ws["A2"].font = sub_font
    cols = [clean(str(c)) for c in df.columns]
    ws.append([])
    ws.append(cols)
    for j in range(1, len(cols) + 1):
        c = ws.cell(row=4, column=j); c.font = hdr_font; c.fill = hdr_fill; c.alignment = center
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                   (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in row])
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14 if j > 1 else 13
    for row in ws.iter_rows(min_row=5, min_col=1, max_col=1):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    if fmt:
        for row in ws.iter_rows(min_row=5, min_col=2, max_col=len(cols)):
            for c in row:
                c.number_format = fmt
    ws.freeze_panes = "B5"
    return ws


def rd(name, **kw):
    return pd.read_csv(f"{R}/{name}.csv", **kw)


full = rd("J_performance_full"); oos = rd("J_performance_oos"); ins = rd("J_performance_is")
daily = rd("J_daily_returns").rename(columns={"Unnamed: 0": "Date"})
daily["Date"] = pd.to_datetime(daily["Date"])
is_cmb = lambda r: str(r.iloc[0]).startswith("Combined portfolio") and "20bps" not in str(r.iloc[0])

# =========================================================================== #
# Overview
# =========================================================================== #
ws = wb.active; ws.title = "Overview"
ws["A1"] = "Statistical Arbitrage in Cryptocurrencies - Data Workbook"; ws["A1"].font = title_font
ws["A2"] = clean(f"Wall Street Quants Bootcamp Project | live data {S['sample_start']} to {S['sample_end']} | "
                 f"backtest {S['bt_start']} to {S['sample_end']}"); ws["A2"].font = sub_font
c = full.set_index("Strategy").loc["Combined portfolio"]
co = oos.set_index("Strategy").loc["Combined portfolio"]
info = [
    ("Data sources", "Coinbase Exchange API (daily + hourly OHLCV); Deribit DVOL (BTC implied vol)"),
    ("Universe", "16 coins: BTC ETH SOL ADA DOGE AVAX LINK LTC DOT ATOM UNI BCH XLM ALGO AAVE COMP"),
    ("Research period", f"{S['bt_start']} to {S['is_end']} (all parameters chosen here)"),
    ("Hold-out period", f"{S['oos_start']} to {S['sample_end']} (never used for selection)"),
    ("Costs", "20 bps per unit turnover for market orders; 7 bps for limit orders"),
    ("Momentum sleeve", "Ensemble: " + "; ".join(S["momentum_components"]) + f" | execution: {S['execution_selected']['Momentum']}"),
    ("Reversal sleeve", f"{S['reversal_selected']} | execution: {S['execution_selected']['Reversal']}"),
    ("Weighting", S["weighting_selected"]),
    ("Candidates tested", str(S["n_candidates"])),
]
r = 4
for k, v in info:
    ws.cell(row=r, column=1, value=k).font = bold_font
    ws.cell(row=r, column=2, value=clean(v)).font = base_font
    r += 1
r += 1
ws.cell(row=r, column=1, value="Combined portfolio").font = Font(name=FONT, bold=True, size=12, color=NAVY)
r += 1
for j, h in enumerate(["Metric", "Full backtest", "Hold-out"], 1):
    cc = ws.cell(row=r, column=j, value=h); cc.font = hdr_font; cc.fill = hdr_fill; cc.alignment = center
r += 1
for lab, key, f in [("Annual return", "Ann.Return", "0.0%"), ("Annual volatility", "Ann.Vol", "0.0%"),
                    ("Sharpe ratio", "Sharpe", "0.00"), ("Max drawdown", "MaxDD", "0.0%"),
                    ("Alpha (annual, vs BTC)", "Alpha", "0.0%"), ("Beta (vs BTC)", "Beta", "0.00"),
                    ("Alpha t-stat", "Alpha t", "0.00")]:
    ws.cell(row=r, column=1, value=lab).font = base_font
    for j, src in [(2, c), (3, co)]:
        cc = ws.cell(row=r, column=j, value=float(src[key])); cc.number_format = f; cc.font = base_font
        cc.alignment = Alignment(horizontal="center")
    r += 1
r += 1
ws.cell(row=r, column=1, value="Sheet index").font = Font(name=FONT, bold=True, size=12, color=NAVY)
r += 1
index_rows = [
    ("Performance", "Returns, vol, Sharpe, max DD, alpha / beta for full, research and hold-out periods"),
    ("Daily Returns", "Daily net returns of every final strategy and benchmark"),
    ("Equity & Drawdown", "Growth of $1 and drawdown series"),
    ("A Horizon Scan", "Momentum vs reversal by look-back (daily)"),
    ("A2 Intraday Horizon", "Momentum vs reversal by look-back (hourly)"),
    ("B Activity", "Volume-shock conditioning"),
    ("C Weekday-Weekend / C Sessions / C Weekend Mom", "Seasonality"),
    ("D Themes", "Theme momentum grid"),
    ("E Front-run US", "Technical play: front-running the US session"),
    ("F Correlation", "Residual reversal and cointegrated pairs"),
    ("G Macro Regimes / G Macro Indicators", "Strategy Sharpe by volatility / dislocation regime"),
    ("H Candidates / H Execution / H Cost Sweep", "Selection grid, execution variants, 0-200 bps sweep"),
    ("I Weighting / I Robustness / I Correlation", "Combining momentum and reversal"),
    ("Pairs Frequency / Pairs Walk-Forward / Pairs Daily", "Cointegration engine outputs"),
    ("Raw Prices / DVOL", "Input data (daily)"),
]
for k, v in index_rows:
    ws.cell(row=r, column=1, value=k).font = bold_font
    ws.cell(row=r, column=2, value=v).font = base_font
    r += 1
ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 110
ws.column_dimensions["C"].width = 14

# =========================================================================== #
# Performance
# =========================================================================== #
perf = pd.concat([full.assign(Period="Full backtest"), ins.assign(Period="Research (IS)"),
                  oos.assign(Period="Hold-out (OOS)")])
perf = perf[["Period"] + [x for x in perf.columns if x != "Period"]]
sheet("Performance", perf, "Performance evaluation",
      "Net of costs. Alpha / beta from daily regression on BTC with Newey-West errors. 365-day year.",
      first_w=16, widths={2: 40}, highlight=lambda r_: str(r_.iloc[1]) == "Combined portfolio")
ws = wb["Performance"]
for row in ws.iter_rows(min_row=5, min_col=2, max_col=2):
    for cc in row:
        cc.alignment = left

big_sheet("Daily Returns", daily, "Daily net returns", "Final strategies and benchmarks, simple daily returns.")
eq = daily.set_index("Date").fillna(0)
wealth = (1 + eq).cumprod()
dd = wealth / wealth.cummax() - 1
ed = pd.concat([wealth.add_prefix("Growth: "), dd[["Combined portfolio", "Buy & Hold BTC"]].add_prefix("Drawdown: ")],
               axis=1).reset_index()
big_sheet("Equity & Drawdown", ed, "Growth of $1 and drawdowns", "Computed from the Daily Returns sheet.", fmt="0.0000")

# =========================================================================== #
# Research blocks
# =========================================================================== #
sheet("A Horizon Scan", rd("A_horizon_scan"), "A. Time horizon (daily)",
      "Momentum-sign strategies by look-back. Positive = momentum, negative = reversal. IS = research, OOS = hold-out.")
sheet("A2 Intraday Horizon", rd("A2_intraday_horizon"), "A2. Time horizon (hourly)",
      "Cross-sectional signal rebalanced hourly, 2022 onward. Negative Sharpe = reversal.", first_w=16)
sheet("B Activity", rd("B_activity_volume"), "B. New information / activity",
      "Cross-sectional momentum conditioned on abnormal dollar volume.", first_w=46)
sheet("C Weekday-Weekend", rd("C_weekday_weekend"), "C. Seasonality - weekdays vs weekends",
      "Gross strategy returns by day type and day of week.", first_w=28)
sheet("C Sessions", rd("C_session_returns"), "C. Seasonality - trading sessions (hourly)",
      "Equal-weight crypto and BTC mean return per hour by session (UTC), 2022 onward.", first_w=14,
      widths={2: 22})
sheet("C Weekend Mom", rd("C_weekend_momentum"), "C. Trading the weekend season",
      "Time-series momentum held all days vs weekends only vs weekdays only.", first_w=44)
sheet("D Themes", rd("D_theme_momentum"), "D. Investment themes", "Theme momentum by look-back and holding period.")
sheet("E Front-run US", rd("E_frontrun_us_session"), "E. Technical plays - front-running the US session",
      "Session strategies open and close daily (turnover 2 per day).", first_w=46)
sheet("F Correlation", rd("F_correlation_reversal"), "F. Correlation-based reversal",
      "Residual reversal vs market and vs theme; cointegrated pairs (Kalman, walk-forward).", first_w=44)
sheet("G Macro Regimes", rd("G_macro_regimes"), "G. Macro regimes",
      "Gross Sharpe by trailing-year tercile of each indicator, known before the trade.", first_w=28,
      widths={2: 32})
mi = rd("G_macro_indicators").rename(columns={"Unnamed: 0": "Date"}); mi["Date"] = pd.to_datetime(mi["Date"])
big_sheet("G Macro Indicators", mi, "G. Macro indicators (daily)",
          "Implied vol (DVOL), realized vol, dispersion, average pairwise correlation.", fmt="0.0000")
sel = rd("H_candidate_selection")
sheet("H Candidates", sel, "H. Candidate selection",
      "Ranked by research-period net Sharpe at the family's execution cost (momentum 20 bps, reversal 7 bps).",
      first_w=46, highlight=lambda r_: r_.iloc[0] in S["momentum_components"] + [S["reversal_selected"]])
sheet("H Execution", rd("H_execution_variants"), "H. Execution variants",
      "Holding periods and no-trade bands for the final books.", first_w=14, widths={2: 24})
sheet("H Cost Sweep", rd("H_cost_sweep"), "H. Transaction-cost sweep",
      "Net Sharpe at 0-200 bps per unit turnover.", first_w=12, fmt_override={"cost_bps": "0"})
sheet("I Weighting", rd("I_weighting_methods"), "I. Weighting methods",
      "Walk-forward weights from the trailing 180 days, rebalanced monthly.", first_w=26,
      highlight=lambda r_: r_.iloc[0] == S["weighting_selected"])
sheet("I Robustness", rd("I_selection_robustness"), "I. Selection robustness",
      "Combined portfolio under alternative momentum sleeves.", first_w=56)
sheet("I Correlation", rd("I_sleeve_correlation").rename(columns={"Unnamed: 0": "Sleeve"}),
      "I. Correlation of daily net returns", "Final sleeves, momentum components, pairs and BTC.", first_w=24)

# =========================================================================== #
# Pairs engine
# =========================================================================== #
freq = pd.read_csv("data/pair_frequency.csv").rename(columns={"Unnamed: 0": "Pair"})
sheet("Pairs Frequency", freq, "Cointegration pairs - selection frequency",
      "Number of monthly formation windows in which each pair passed all filters.", first_w=16)
sellog = pd.read_csv("data/selection_log.csv")
sheet("Pairs Walk-Forward", sellog, "Cointegration pairs - walk-forward log",
      "Pairs selected at the end of each 180-day formation window.", first_w=16, widths={3: 120})
cd = pd.read_csv("data/coint_daily.csv").rename(columns={"Unnamed: 0": "Date"}); cd["Date"] = pd.to_datetime(cd["Date"])
big_sheet("Pairs Daily", cd, "Cointegration pairs - daily P&L", "Gross, net (20 bps), turnover, active pairs.",
          fmt="0.0000")

# =========================================================================== #
# Inputs
# =========================================================================== #
raw = pd.read_csv("data/crypto_raw_data.csv", parse_dates=["date"])
big_sheet("Raw Prices", raw[["date", "asset", "open", "high", "low", "close", "volume"]].rename(columns=str.title),
          "Daily OHLCV (Coinbase Exchange)", f"{raw['date'].min().date()} to {raw['date'].max().date()}; "
          "includes XRP and MKR, which are excluded from the traded universe.", fmt="#,##0.0000")
ws = wb["Raw Prices"]
for row in ws.iter_rows(min_row=5, min_col=2, max_col=2):
    for cc in row:
        cc.number_format = "General"
dv = pd.read_csv("data/dvol.csv").rename(columns={"Unnamed: 0": "Date", "date": "Date"})
dv["Date"] = pd.to_datetime(dv["Date"])
big_sheet("DVOL", dv, "Deribit DVOL implied volatility index", "Daily close, BTC and ETH.", fmt="0.00")

wb.save("crypto_stat_arb_data.xlsx")
print("saved crypto_stat_arb_data.xlsx with", len(wb.sheetnames), "sheets")
